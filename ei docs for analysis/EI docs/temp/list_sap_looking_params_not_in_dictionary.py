"""
Scan Parts1_4_enriched Part 1-4 DOCX (Parameter Configuration Guidelines + param heading regex).

Build set of parameters that *look* SAP-standard:
  - Union: curated SAP-positive lists (unified, OVERRIDES, attached canonical, twice overlay CANONICAL minus NON_SAP)
  - Plus: DDIC-style names matching ^[A-Z][A-Z0-9_]{2,29}$ after excluding obvious composite/UI/reporting patterns

Subtract params_dictionary.xlsx column A.

Writes: sap_standard_looking_not_in_params_dictionary.csv
"""

from __future__ import annotations

import csv
import importlib.util
import re
from collections import Counter
from pathlib import Path

from docx import Document
import openpyxl

ROOT = Path(r"c:\vibe code dev\ei_auto_doc")
TEMP = ROOT / "ei docs for analysis" / "EI docs" / "temp"
ENRICHED = TEMP / "Parts1_4_enriched"
MAIN_DICT = TEMP / "params_dictionary.xlsx"
OUT_CSV = TEMP / "sap_standard_looking_not_in_params_dictionary.csv"

RE_PARAM_HEAD = re.compile(r"^([A-Z][A-Z0-9_]*)\s*\(.*\)\s*:?\s*$")
RE_SAP_SHAPE = re.compile(r"^[A-Z][A-Z0-9_]{2,29}$")
STOP_TITLES = {
    "parameter relationship",
    "parameter relationships",
    "default values",
    "practical example of parameter configuration",
    "practical configuration examples",
    "ei function structure",
    "abap code",
}

# Obvious composite / presentation / typo patterns — not counted as "SAP-looking" by shape alone.
EXCLUDE_PREFIXES = (
    "BALANCE_",
    "NETWR_",
    "FILE_SIZE",
    "PARKED_",
    "SEVE_",
    "SEVER",
    "PROCSTAT_",
    "PLANT_",
    "SHIPTO_",
    "GL_ACC_",
    "ERNAM_",
    "TIME_DIFF",
    "PACKAGE_",
    "IP_ADDRESS",
    "PROCESSID",
    "ERRORNAME",
    "ERR_MSG",
    "USES_32",
    "US_GUI",
    "US_PLUGIN",
    "US_RFC",
    "VALUEORIG",
    "SLGCONNECTION",
    "SLGDATA",
    "SLGPASSPORT",
    "SLGROOT",
    "SLGTERM",
    "TOTAL_MEM",
    "POSSIBLE_",
    "PRESENT_",
    "SEVERFLTRD",
)

EXCLUDE_SUFFIXES = (
    "_DESC",
    "_TEXT",
    "_TXT",
    "_FR",
    "_KB",
    "_MB",
    "_LIPS",
    "_VBAP",
    "_IVNOICE",
)

EXCLUDE_EXACT = frozenset(
    {
        "AMOUNT",
        "ACTIVITY",
        "COUNTER",
        "MESSAGE",
        "MSG",
        "PERIOD",
        "SERVER",
        "RFC",
        "PROFILE",
        "PROJECT",
        "PROTOCOL",
        "TRACE",
        "OWNER",
        "MASTER",
        "COUNT",
        "TYPE",
        "STATUS",
        "TERM",
        "ROUTE",
        "ROLL",
    }
)


def iter_param_heads(fp: Path):
    doc = Document(str(fp))
    in_g = False
    for p in doc.paragraphs:
        t = (p.text or "").strip()
        lo = t.lower()
        if lo == "parameter configuration guidelines":
            in_g = True
            continue
        if in_g and lo in STOP_TITLES:
            in_g = False
            continue
        if not in_g:
            continue
        m = RE_PARAM_HEAD.match(t)
        if m:
            yield m.group(1).upper()


def load_curated_sap_sets() -> tuple[set[str], set[str], set[str], set[str], frozenset[str]]:
    spec = importlib.util.spec_from_file_location("u", ROOT / "tools" / "sap_unified_param_texts.py")
    m = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(m)
    uni = {k.upper() for k in m.SAP_UNIFIED_EXPLANATION}

    spec2 = importlib.util.spec_from_file_location("b", TEMP / "build_params_dictionary_xlsx.py")
    m2 = importlib.util.module_from_spec(spec2)
    assert spec2.loader
    spec2.loader.exec_module(m2)
    ovr = {k.upper() for k in m2.OVERRIDES}

    spec3 = importlib.util.spec_from_file_location("a", TEMP / "sap_canonical_attached_params.py")
    m3 = importlib.util.module_from_spec(spec3)
    assert spec3.loader
    spec3.loader.exec_module(m3)
    att = {k.upper() for k in m3.SAP_CANONICAL_ATTACHED}

    spec4 = importlib.util.spec_from_file_location("o", TEMP / "sap_twice_csv_overlay.py")
    m4 = importlib.util.module_from_spec(spec4)
    assert spec4.loader
    spec4.loader.exec_module(m4)
    non_sap = frozenset(str(x).upper() for x in m4.NON_SAP_STANDARD)
    canon = {k.upper() for k in m4.CANONICAL if k.upper() not in non_sap}

    return uni, ovr, att, canon, non_sap


def shape_heuristic(p: str) -> bool:
    if not RE_SAP_SHAPE.match(p):
        return False
    if p in EXCLUDE_EXACT:
        return False
    for pre in EXCLUDE_PREFIXES:
        if p.startswith(pre):
            return False
    for suf in EXCLUDE_SUFFIXES:
        if p.endswith(suf):
            return False
    return True


def is_sap_standard_looking(
    p: str,
    uni: set[str],
    ovr: set[str],
    att: set[str],
    canon: set[str],
    non_sap: frozenset[str],
) -> bool:
    if p in non_sap:
        return False
    if p in uni or p in ovr or p in att or p in canon:
        return True
    return shape_heuristic(p)


def load_dictionary_keys() -> set[str]:
    wb = openpyxl.load_workbook(MAIN_DICT, read_only=True, data_only=True)
    ws = wb["dictionary"]
    s = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        v = row[0]
        if v is None or str(v).strip() == "":
            continue
        s.add(str(v).strip().upper())
    wb.close()
    return s


def main() -> None:
    uni, ovr, att, canon, non_sap = load_curated_sap_sets()
    dict_keys = load_dictionary_keys()

    ctr: Counter[str] = Counter()
    for part in ("Part 1", "Part 2", "Part 3", "Part 4"):
        d = ENRICHED / part
        if not d.is_dir():
            continue
        for fp in sorted(p for p in d.glob("*.docx") if not p.name.startswith("~$")):
            for pname in iter_param_heads(fp):
                ctr[pname] += 1

    corpus = set(ctr.keys())
    sap_looking = {p for p in corpus if is_sap_standard_looking(p, uni, ovr, att, canon, non_sap)}
    not_covered = sorted(sap_looking - dict_keys)

    with OUT_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["parameter", "total_hits_guidelines"])
        for p in not_covered:
            w.writerow([p, ctr[p]])

    print(f"Corpus unique params (guidelines): {len(corpus)}")
    print(f"SAP-standard-looking (heuristic + curated): {len(sap_looking)}")
    print(f"In params_dictionary.xlsx: {len(dict_keys)}")
    print(f"SAP-looking NOT in dictionary: {len(not_covered)}")
    print(f"Wrote: {OUT_CSV}")


if __name__ == "__main__":
    main()
